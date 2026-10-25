"""Tests voor tools/protocollen.py. Uitvoeren met:  tools/.venv/bin/python -m unittest discover -s tools -v

Elke test werkt op een kopie van de protocolbestanden en sw.js in een tijdelijke map: de echte bestanden
blijven onaangeroerd."""

import contextlib
import datetime
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import protocollen as tool  # noqa: E402

REAL = Path(__file__).resolve().parent.parent


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.data = self.tmp / "data"
        shutil.copytree(REAL / "data", self.data)
        shutil.copy(REAL / "sw.js", self.tmp / "sw.js")
        self.xlsx = self.tmp / "protocollen.xlsx"
        self.today = datetime.date.today().isoformat()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def run_tool(self, *args):
        out = io.StringIO()
        err = io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = tool.main(["--data", str(self.data), "--sw", str(self.tmp / "sw.js"), *args])
        return code, out.getvalue() + err.getvalue()

    def export(self):
        code, text = self.run_tool("export", "--uit", str(self.xlsx))
        self.assertEqual(code, 0, text)

    def importing(self, apply=False):
        args = ["import", "--van", str(self.xlsx)] + (["--toepassen"] if apply else [])
        return self.run_tool(*args)

    def protocol(self, name="vkb"):
        return json.loads((self.data / f"protocol-{name}.json").read_text(encoding="utf-8"))

    def cache_version(self):
        import re
        return int(re.search(r"CACHE_VERSION = 'v(\d+)'", (self.tmp / "sw.js").read_text()).group(1))

    def cell(self, sheet, ident, header, value=None, id_header="ID", read=False):
        """Zoek de cel van een rij (op id) en kolomtitel; schrijf er een waarde in, of lees ze."""
        wb = load_workbook(self.xlsx)
        ws = wb[sheet]
        headers = {ws.cell(row=3, column=c).value: c for c in range(1, ws.max_column + 1)}
        for r in range(4, ws.max_row + 1):
            if str(ws.cell(row=r, column=headers[id_header]).value) == str(ident):
                target = ws.cell(row=r, column=headers[header])
                if read:
                    return target.value
                target.value = value
                wb.save(self.xlsx)
                return None
        raise AssertionError(f"rij {ident} niet gevonden in {sheet}")

    def criterion(self, cid, name="vkb"):
        return next(c for lv in self.protocol(name)["levels"] for c in lv["criteria"] if c["id"] == cid)


class Opbouw(Base):
    def test_werkboek_heeft_de_bladen_een_verborgen_kopie_en_bescherming(self):
        self.export()
        wb = load_workbook(self.xlsx)
        self.assertEqual(wb.sheetnames, ["Leesmij", "Testschema", "Referenties", "Criteria", "Levels", "Tests", "Mijlpalen", "Stopregels", "Instellingen", "_origineel"])
        self.assertEqual(wb["_origineel"].sheet_state, "hidden")
        self.assertTrue(all(wb[name].protection.sheet for name in wb.sheetnames if name != "_origineel"))

    def test_enkel_gele_cellen_zijn_ontgrendeld(self):
        self.export()
        ws = load_workbook(self.xlsx)["Criteria"]
        unlocked = {ws.cell(row=3, column=c).value for r in range(4, ws.max_row + 1) for c in range(1, ws.max_column + 1) if not ws.cell(row=r, column=c).protection.locked}
        self.assertEqual(unlocked, {"Omschrijving", "Tekst voor de patiënt", "Grens", "Niveau (A–D)", "Testmethode", "Bij niet halen", "Notitie"})

    def test_grens_en_omschrijving_van_een_criterium_staan_in_het_werkboek(self):
        self.export()
        self.assertEqual(self.cell("Criteria", "vkb-6-quad", "Grens", read=True), 80)
        self.assertIn("{threshold}", self.cell("Criteria", "vkb-6-quad", "Omschrijving", read=True))
        formula = self.cell("Criteria", "vkb-6-quad", "Zo ziet het eruit", read=True)
        self.assertTrue(formula.startswith("=SUBSTITUTE("))

    def test_zonder_wijziging_staat_er_niets_te_verwerken(self):
        self.export()
        code, text = self.importing()
        self.assertEqual(code, 0)
        self.assertIn("Geen wijzigingen", text)


class Wijzigen(Base):
    def test_grens_aanpassen_eerst_controleren_dan_toepassen(self):
        self.export()
        sw_before = (self.tmp / "sw.js").read_text()
        v_before = self.cache_version()
        enkel_before = (self.data / "protocol-enkel.json").read_bytes()
        version_before = self.protocol()["version"]
        self.cell("Criteria", "vkb-6-quad", "Grens", 75)

        code, text = self.importing()  # alleen controleren
        self.assertEqual(code, 0, text)
        self.assertIn("vkb-6-quad", text)
        self.assertIn("80 → 75", text)
        self.assertEqual(self.criterion("vkb-6-quad")["threshold"], 80, "dry-run verandert niets")
        self.assertEqual((self.tmp / "sw.js").read_text(), sw_before)

        code, text = self.importing(apply=True)
        self.assertEqual(code, 0, text)
        crit = self.criterion("vkb-6-quad")
        self.assertEqual(crit["threshold"], 75)
        self.assertIn("{threshold}", crit["label"], "de omschrijving blijft een sjabloon")
        self.assertEqual(self.protocol()["version"], tool.bump_version(version_before, self.today))
        self.assertNotEqual(self.protocol()["version"], version_before, "de versie van het protocol is verhoogd")
        self.assertEqual(self.cache_version(), v_before + 1)
        self.assertEqual((self.data / "protocol-enkel.json").read_bytes(), enkel_before, "het enkelprotocol blijft onaangeroerd")

    def test_aanpassen_van_de_enkel_laat_vkb_ongemoeid(self):
        self.export()
        vkb_before = (self.data / "protocol-vkb.json").read_bytes()
        self.cell("Criteria", "enk-4-vragenlijst", "Grens", 24)
        self.assertEqual(self.importing(apply=True)[0], 0)
        self.assertEqual(self.criterion("enk-4-vragenlijst", "enkel")["threshold"], 24)
        self.assertEqual((self.data / "protocol-vkb.json").read_bytes(), vkb_before)

    def test_een_termijn_in_weken_en_in_maanden(self):
        self.export()
        self.cell("Criteria", "vkb-2-tijd", "Grens", 14)
        self.cell("Criteria", "vkb-4-tijd", "Grens", 10)
        self.assertEqual(self.importing(apply=True)[0], 0)
        self.assertEqual(self.criterion("vkb-2-tijd")["weeks"], 14)
        self.assertEqual(self.criterion("vkb-4-tijd")["months"], 10)

    def test_komma_als_decimaalteken_en_tekst_met_een_getal(self):
        self.export()
        self.cell("Criteria", "enk-2-wblt", "Grens", "2,5")
        self.assertEqual(self.importing(apply=True)[0], 0)
        self.assertEqual(self.criterion("enk-2-wblt", "enkel")["threshold"], 2.5)

    def test_teksten_en_regels_per_zin(self):
        self.export()
        self.cell("Levels", 2, "Doel voor de patiënt", "Nieuw doel.", id_header="Level")
        self.cell("Levels", 2, "Wat doen we nu? (patiënt, één regel per zin)", "Eerste zin.\n\n  Tweede zin.  \nDerde zin.", id_header="Level")
        self.assertEqual(self.importing(apply=True)[0], 0)
        level = next(lv for lv in self.protocol()["levels"] if lv["number"] == 2)
        self.assertEqual(level["patientGoal"], "Nieuw doel.")
        self.assertEqual(level["patientFocus"], ["Eerste zin.", "Tweede zin.", "Derde zin."])

    def test_een_optionele_cel_leegmaken_haalt_de_waarde_weg(self):
        self.export()
        self.assertIn("ifNotMet", self.criterion("vkb-6-quad"))
        self.cell("Criteria", "vkb-6-quad", "Bij niet halen", None)
        self.cell("Tests", "quad-iso", "Doelzone (%)", None)
        self.assertEqual(self.importing(apply=True)[0], 0)
        self.assertNotIn("ifNotMet", self.criterion("vkb-6-quad"))
        test = next(t for g in self.protocol()["testGroups"] for t in g["tests"] if t["id"] == "quad-iso")
        self.assertNotIn("target", test)

    def test_niveau_en_stopregel_aanpassen_en_een_ongeldig_niveau_weigeren(self):
        self.export()
        self.cell("Criteria", "vkb-6-quad", "Niveau (A–D)", "B/C")
        self.cell("Stopregels", 1, "Actie", "Nieuwe actie.", id_header="Nr")
        self.assertEqual(self.importing(apply=True)[0], 0)
        self.assertEqual(self.criterion("vkb-6-quad")["evidence"], "B/C")
        self.assertEqual(self.protocol()["stopRules"][0]["action"], "Nieuwe actie.")
        self.export()
        self.cell("Criteria", "vkb-6-quad", "Niveau (A–D)", "E")
        code, text = self.importing(apply=True)
        self.assertEqual(code, 1, text)
        self.assertIn("A, B, C of D", text)

    def test_testmethode_en_opmerkingen_van_een_fase_aanpassen(self):
        self.export()
        self.cell("Criteria", "vkb-6-quad", "Testmethode", "Handheld dynamometer")
        self.cell("Levels", 2, "Opmerkingen en aandachtspunten (voor jou, één per regel)", "Eerste.\nTweede.", id_header="Level")
        self.assertEqual(self.importing(apply=True)[0], 0)
        self.assertEqual(self.criterion("vkb-6-quad")["method"], "Handheld dynamometer")
        level = next(lv for lv in self.protocol()["levels"] if lv["number"] == 2)
        self.assertEqual(level["remarks"], ["Eerste.", "Tweede."])

    def test_een_referentiewaarde_aanpassen_voor_man_en_vrouw(self):
        self.export()
        self.cell("Referenties", "ref-cmj-hoogte-performance-tracker", "Referentie mannen", "18-26 cm")
        self.cell("Referenties", "ref-pijn", "Referentie (voor iedereen)", "0-2/10")
        self.assertEqual(self.importing(apply=True)[0], 0)
        refs = {r["id"]: r for r in self.protocol()["references"]}
        self.assertEqual(refs["ref-cmj-hoogte-performance-tracker"]["men"], "18-26 cm")
        self.assertEqual(refs["ref-cmj-hoogte-performance-tracker"]["women"], "10-18 cm")
        self.assertEqual(refs["ref-pijn"]["text"], "0-2/10")

    def test_mijlpaalgrens_en_instelling(self):
        self.export()
        self.cell("Mijlpalen", "vkb-lsi80", "Grens", 85)
        self.cell("Mijlpalen", "enk-preventie", "Grens", 8)
        self.cell("Instellingen", "Rustdag-jokers per week", "Waarde", 2, id_header="Instelling")
        self.assertEqual(self.importing(apply=True)[0], 0)
        self.assertEqual(next(m for m in self.protocol()["milestones"] if m["id"] == "vkb-lsi80")["trigger"]["threshold"], 85)
        self.assertEqual(next(m for m in self.protocol("enkel")["milestones"] if m["id"] == "enk-preventie")["trigger"]["weeks"], 8)
        self.assertEqual(self.protocol()["homeLog"]["jokersPerWeek"], 2)

    def test_eenheden_nm_en_n_hebben_elk_een_eigen_bereik(self):
        self.export()
        wb = load_workbook(self.xlsx)
        ws = wb["Tests"]
        headers = {ws.cell(row=3, column=c).value: c for c in range(1, ws.max_column + 1)}
        rows = {(ws.cell(row=r, column=headers["ID"]).value, ws.cell(row=r, column=headers["Eenheid"]).value): r for r in range(4, ws.max_row + 1)}
        ws.cell(row=rows[("quad-iso", "N")], column=headers["Max"]).value = 3000
        ws.cell(row=rows[("quad-iso", "Nm")], column=headers["Max"]).value = 1200
        wb.save(self.xlsx)
        self.assertEqual(self.importing(apply=True)[0], 0)
        test = next(t for g in self.protocol()["testGroups"] for t in g["tests"] if t["id"] == "quad-iso")
        self.assertEqual([(u["unit"], u["max"]) for u in test["units"]], [("Nm", 1200), ("N", 3000)])
        self.assertEqual(test["max"], 1200, "test.max volgt de standaardeenheid")

    def test_testschema_een_moment_toevoegen_en_een_label_wijzigen(self):
        self.export()
        wb = load_workbook(self.xlsx)
        ws = wb["Testschema"]
        last = next(r for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=2).value is None)
        ws.cell(row=last, column=2).value = 130
        ws.cell(row=last, column=4).value = "Afsluitende meting"
        ws.cell(row=5, column=4).value = "Eerste maximale meting (aangepast)"
        wb.save(self.xlsx)
        code, text = self.importing(apply=True)
        self.assertEqual(code, 0, text)
        self.assertIn("toegevoegd: week 130: Afsluitende meting", text)
        self.assertNotIn("week 48: Opvolgmeting; ", text, "geen volledige lijst in de uitvoer")
        moments = self.protocol()["testSchedule"]["moments"]
        self.assertEqual(moments[-1], {"week": 130, "label": "Afsluitende meting"})
        self.assertEqual(tool.moment_changes([(8, 10, "a"), (12, None, "b")], [(8, 10, "a2"), (52, None, "c")]),
                         ["week 8-10: 'a' → 'a2'", "verwijderd: week 12: b", "toegevoegd: week 52: c"])
        self.assertEqual(moments[1], {"week": 8, "weekTo": 10, "label": "Eerste maximale meting (aangepast)"})
        self.assertEqual(len(moments), 17)

    def test_versie_krijgt_een_datum_en_op_dezelfde_dag_een_volgnummer(self):
        self.assertEqual(tool.bump_version("2026-10-01", "2026-10-24"), "2026-10-24")
        self.assertEqual(tool.bump_version("2026-10-24", "2026-10-24"), "2026-10-24.2")
        self.assertEqual(tool.bump_version("2026-10-24.2", "2026-10-24"), "2026-10-24.3")
        self.assertEqual(tool.bump_version("", "2026-10-24"), "2026-10-24")
        # in de praktijk: twee keer na elkaar toepassen verhoogt telkens
        self.export()
        self.cell("Criteria", "vkb-6-quad", "Grens", 75)
        self.importing(apply=True)
        first = self.protocol()["version"]
        self.export()
        self.cell("Criteria", "vkb-6-quad", "Grens", 70)
        self.importing(apply=True)
        self.assertEqual(self.protocol()["version"], tool.bump_version(first, self.today))
        self.assertNotEqual(self.protocol()["version"], first)


class Veiligheid(Base):
    def assert_geweigerd(self, fragment):
        before = {p: (self.data / f"protocol-{p}.json").read_bytes() for p in tool.PROTOCOLS}
        sw = (self.tmp / "sw.js").read_text()
        code, text = self.importing(apply=True)
        self.assertEqual(code, 1, text)
        self.assertIn(fragment, text)
        self.assertEqual({p: (self.data / f"protocol-{p}.json").read_bytes() for p in tool.PROTOCOLS}, before, "niets is veranderd")
        self.assertEqual((self.tmp / "sw.js").read_text(), sw)

    def test_geen_getal_in_een_getalcel(self):
        self.export()
        self.cell("Criteria", "vkb-6-quad", "Grens", "veel")
        self.assert_geweigerd("geen geldig getal")

    def test_lsi_boven_150_procent(self):
        self.export()
        self.cell("Criteria", "vkb-6-quad", "Grens", 200)
        self.assert_geweigerd("moet tussen 0 en 150")

    def test_een_termijn_moet_een_geheel_getal_zijn(self):
        self.export()
        self.cell("Criteria", "vkb-2-tijd", "Grens", 12.5)
        self.assert_geweigerd("geen geldig getal")

    def test_min_moet_kleiner_zijn_dan_max(self):
        self.export()
        self.cell("Tests", "nprs-adl", "Min", 20)
        self.assert_geweigerd("'Min' moet kleiner zijn dan 'Max'")

    def test_een_verplichte_tekst_mag_niet_leeg(self):
        self.export()
        self.cell("Criteria", "vkb-1-gang", "Omschrijving", None)
        self.assert_geweigerd("mag niet leeg zijn")

    def test_placeholder_zonder_getal_wordt_geweigerd(self):
        self.export()
        self.cell("Criteria", "vkb-1-gang", "Omschrijving", "Stappen vanaf {threshold}")  # een oordeel-criterium heeft geen grens
        self.assert_geweigerd("{threshold} staat in de tekst")

    def test_weken_van_het_testschema_moeten_oplopen(self):
        self.export()
        wb = load_workbook(self.xlsx)
        wb["Testschema"].cell(row=6, column=2).value = 5  # week 12 wordt 5, na week 8
        wb.save(self.xlsx)
        self.assert_geweigerd("moeten oplopen")

    def test_conflict_als_het_protocol_intussen_veranderd_is(self):
        self.export()
        path = self.data / "protocol-vkb.json"
        d = self.protocol()
        for lv in d["levels"]:
            for c in lv["criteria"]:
                if c["id"] == "vkb-6-quad":
                    c["threshold"] = 82  # iemand anders (bv. Claude) paste dit aan na de export
        path.write_text(json.dumps(d, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        self.cell("Criteria", "vkb-6-quad", "Grens", 75)
        self.assert_geweigerd("in het protocol staat nu 82")

    def test_andere_wijzigingen_in_het_protocol_geven_geen_conflict(self):
        self.export()
        d = self.protocol()
        d["levels"][0]["title"] = "Nieuwe titel intussen"  # een ander veld dan wat je in Excel aanpast
        (self.data / "protocol-vkb.json").write_text(json.dumps(d, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        self.cell("Criteria", "vkb-6-quad", "Grens", 75)
        self.assertEqual(self.importing(apply=True)[0], 0)
        self.assertEqual(self.protocol()["levels"][0]["title"], "Nieuwe titel intussen", "de nieuwere wijziging blijft staan")
        self.assertEqual(self.criterion("vkb-6-quad")["threshold"], 75)

    def test_opnieuw_inlezen_van_een_reeds_doorgevoerd_bestand_is_geen_conflict(self):
        self.export()
        self.cell("Criteria", "vkb-6-quad", "Grens", 75)
        self.assertEqual(self.importing(apply=True)[0], 0)
        code, text = self.importing()
        self.assertEqual(code, 0, text)
        self.assertIn("Geen wijzigingen", text)

    def test_een_waarschuwing_als_de_omschrijving_nog_het_oude_getal_bevat(self):
        self.export()
        self.cell("Criteria", "vkb-6-quad", "Omschrijving", "Quadriceps LSI ≥ 80 %")  # vast getal in plaats van {threshold}
        self.cell("Criteria", "vkb-6-quad", "Grens", 75)
        code, text = self.importing()
        self.assertEqual(code, 0, text)
        self.assertIn("LET OP", text)
        self.assertIn("{threshold}", text)

    def test_een_bestand_van_een_andere_bron_of_met_gewijzigde_kolomtitel(self):
        self.export()
        wb = load_workbook(self.xlsx)
        wb["Criteria"].cell(row=3, column=11).value = "Limiet"
        wb.save(self.xlsx)
        code, text = self.importing()
        self.assertEqual(code, 1)
        self.assertIn("kolomtitel", text)

        wb = load_workbook(self.xlsx)
        del wb["_origineel"]
        wb.save(self.xlsx)
        code, text = self.importing()
        self.assertEqual(code, 1)
        self.assertIn("verborgen blad ontbreekt", text)

    def test_ontbrekend_bestand(self):
        code, text = self.importing()
        self.assertEqual(code, 1)
        self.assertIn("bestaat niet", text)


if __name__ == "__main__":
    unittest.main()
